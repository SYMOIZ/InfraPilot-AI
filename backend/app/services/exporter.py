"""Export architectures to draw.io, PNG placeholder, and Excel workbooks."""

from __future__ import annotations

import io
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageDraw, ImageFont

from app.models.schemas import Architecture, EvaluationReport, Project, Requirements
from app.models.validation import load_catalog
from app.services.drawio_builder import build_drawio_xml


def export_drawio(architecture: Architecture) -> bytes:
    return build_drawio_xml(architecture).encode("utf-8")


def export_png(architecture: Architecture, width: int = 1000, height: int = 700) -> bytes:
    """Generate a simple placeholder PNG diagram via Pillow (not full draw.io fidelity)."""
    img = Image.new("RGB", (width, height), color=(245, 247, 250))
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.load_default()
        title_font = ImageFont.load_default()
    except Exception:
        font = None
        title_font = None

    draw.rectangle([0, 0, width, 48], fill=(15, 23, 42))
    draw.text((16, 14), f"InfraPilot: {architecture.name}", fill=(255, 255, 255), font=title_font)
    draw.text(
        (16, 56),
        "Placeholder PNG export (Pillow) — open .drawio for editable diagram.",
        fill=(71, 85, 105),
        font=font,
    )

    for idx, node in enumerate(architecture.nodes):
        x = int(node.x) if node.x is not None else 40 + (idx % 4) * 220
        y = int(node.y) if node.y is not None else 100 + (idx // 4) * 140
        x = max(20, min(x, width - 160))
        y = max(90, min(y, height - 80))
        draw.rounded_rectangle(
            [x, y, x + 140, y + 64],
            radius=8,
            fill=(226, 232, 240),
            outline=(51, 65, 85),
            width=2,
        )
        draw.text((x + 10, y + 12), node.label[:22], fill=(15, 23, 42), font=font)
        draw.text((x + 10, y + 36), node.service_id, fill=(71, 85, 105), font=font)

    # Simple edge lines using node positions
    pos = {}
    for idx, node in enumerate(architecture.nodes):
        x = int(node.x) if node.x is not None else 40 + (idx % 4) * 220
        y = int(node.y) if node.y is not None else 100 + (idx // 4) * 140
        pos[node.id] = (x + 70, y + 32)

    for edge in architecture.edges:
        if edge.from_node in pos and edge.to_node in pos:
            draw.line([pos[edge.from_node], pos[edge.to_node]], fill=(100, 116, 139), width=2)

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def export_excel(project: Project) -> bytes:
    """Build multi-worksheet Excel workbook with qualitative cost notes only."""
    wb = Workbook()
    header_font = Font(bold=True)

    architecture = project.architecture
    requirements = project.requirements
    evaluation = project.evaluation
    optimized = project.optimized_architecture or (
        project.optimization.optimized if project.optimization else None
    )
    catalog = {s["id"]: s for s in load_catalog().get("services", [])}

    # 1. Architecture Overview
    ws = wb.active
    ws.title = "Architecture Overview"
    ws.append(["Field", "Value"])
    for cell in ws[1]:
        cell.font = header_font
    ws.append(["Project ID", project.id])
    ws.append(["Project Name", project.name])
    ws.append(["Description", project.description or ""])
    if architecture:
        ws.append(["Architecture Name", architecture.name])
        ws.append(["Source", architecture.source])
        ws.append(["Confirmed", architecture.confirmed])
        ws.append(["Node Count", len(architecture.nodes)])
        ws.append(["Edge Count", len(architecture.edges)])
        ws.append(["Region", architecture.metadata.get("region", "")])
    if requirements:
        ws.append(["Purpose", requirements.purpose])
        ws.append(["Workload", requirements.workload.value])
        ws.append(["Availability", requirements.availability.value])
        ws.append(["Budget (qualitative)", requirements.budget])
    ws.append(
        [
            "Disclaimer",
            "InfraPilot does not deploy infrastructure or generate Terraform/CloudFormation.",
        ]
    )

    # 2. AWS Service Inventory
    ws2 = wb.create_sheet("AWS Service Inventory")
    ws2.append(
        ["Node ID", "Service ID", "Service Name", "Label", "Status", "Props"]
    )
    for cell in ws2[1]:
        cell.font = header_font
    if architecture:
        for node in architecture.nodes:
            name = catalog.get(node.service_id, {}).get("name", node.service_id)
            status = "implemented" if architecture.confirmed else "planned"
            ws2.append(
                [
                    node.id,
                    node.service_id,
                    name,
                    node.label,
                    status,
                    str(node.props),
                ]
            )

    # 3. Data Flow
    ws3 = wb.create_sheet("Data Flow")
    ws3.append(["Edge ID", "From", "To", "Label"])
    for cell in ws3[1]:
        cell.font = header_font
    if architecture:
        for edge in architecture.edges:
            ws3.append([edge.id or "", edge.from_node, edge.to_node, edge.label or ""])

    # 4. Validation Rules
    ws4 = wb.create_sheet("Validation Rules")
    ws4.append(
        ["Rule ID", "Category", "Title", "Status", "Severity", "Message", "Source"]
    )
    for cell in ws4[1]:
        cell.font = header_font
    if evaluation:
        for f in evaluation.findings:
            ws4.append(
                [
                    f.rule_id,
                    f.category,
                    f.title,
                    f.status.value,
                    f.severity,
                    f.message,
                    f.source,
                ]
            )

    # 5. Security and IAM
    ws5 = wb.create_sheet("Security and IAM")
    ws5.append(["Note", "Detail"])
    for cell in ws5[1]:
        cell.font = header_font
    ws5.append(
        [
            "Policy note",
            "Do not invent production IAM policy documents. Document intent only.",
        ]
    )
    if architecture:
        for node in architecture.nodes:
            if node.service_id in ("iam", "kms", "waf", "cognito", "secrets_manager"):
                ws5.append([node.label, str(node.props)])
    if evaluation:
        for f in evaluation.findings:
            if f.category == "security":
                ws5.append([f.rule_id, f"{f.status.value}: {f.message}"])

    # 6. Risks and Recommendations
    ws6 = wb.create_sheet("Risks and Recommendations")
    ws6.append(["Finding", "Status", "Recommendation"])
    for cell in ws6[1]:
        cell.font = header_font
    if evaluation:
        for f in evaluation.findings:
            if f.status.value in ("Fail", "Warning", "Unknown"):
                ws6.append([f.title, f.status.value, f.recommendation or ""])

    # 7. CI/CD and Operations (sheet titles cannot contain '/')
    ws7 = wb.create_sheet("CI-CD and Operations")
    ws7.append(["Topic", "Guidance"])
    for cell in ws7[1]:
        cell.font = header_font
    ws7.append(
        [
            "CI/CD",
            "InfraPilot evaluates architecture only - pipeline tooling is out of scope for deployment generation.",
        ]
    )
    ws7.append(
        [
            "Observability",
            "Prefer CloudWatch metrics, logs, and alarms as documented in the architecture.",
        ]
    )
    if architecture and any(n.service_id == "cloudwatch" for n in architecture.nodes):
        ws7.append(["CloudWatch", "Present in architecture (planned/implemented per confirmation)."])
    else:
        ws7.append(["CloudWatch", "Not present — add monitoring before production."])

    # 8. Cost Assumptions
    ws8 = wb.create_sheet("Cost Assumptions")
    ws8.append(["Assumption", "Detail"])
    for cell in ws8[1]:
        cell.font = header_font
    ws8.append(
        [
            "Pricing",
            "No AWS list prices are invented. Use AWS Pricing Calculator for estimates.",
        ]
    )
    if requirements:
        ws8.append(["Budget (qualitative)", requirements.budget])
        ws8.append(["Traffic (qualitative)", requirements.traffic])
    if project.optimization:
        ws8.append(["Optimization cost notes", project.optimization.cost_notes])
    if optimized:
        ws8.append(
            [
                "Optimized services",
                ", ".join(sorted({n.service_id for n in optimized.nodes})),
            ]
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
