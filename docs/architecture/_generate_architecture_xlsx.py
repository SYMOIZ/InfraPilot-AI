"""Generate InfraPilot_AI_Architecture.xlsx — run from repo root or this folder."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent / "InfraPilot_AI_Architecture.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1565C0")
HEADER_FONT = Font(color="FFFFFF", bold=True)
IMPL_FILL = PatternFill("solid", fgColor="C8E6C9")
PLAN_FILL = PatternFill("solid", fgColor="FFF9C4")
WARN_FILL = PatternFill("solid", fgColor="FFE0B2")


def style_header(ws, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(1, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autosize(ws, min_w: int = 12, max_w: int = 48) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = min_w
        for cell in col:
            if cell.value is not None:
                length = max(length, min(max_w, len(str(cell.value)) + 2))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[letter].width = length


def add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    style_header(ws, len(headers))
    for row in rows:
        ws.append(row)
        status_idx = None
        for i, h in enumerate(headers, 1):
            if h.lower() in ("status", "state", "implementation"):
                status_idx = i
                break
        if status_idx and row:
            val = str(row[status_idx - 1]).lower()
            fill = None
            if "implement" in val:
                fill = IMPL_FILL
            elif "planned" in val or "scaffold" in val or "placeholder" in val:
                fill = PLAN_FILL
            elif "n/a" in val or "out of scope" in val:
                fill = WARN_FILL
            if fill:
                ws.cell(ws.max_row, status_idx).fill = fill
    autosize(ws)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def main() -> None:
    wb = Workbook()
    # Remove default sheet
    default = wb.active
    wb.remove(default)

    add_sheet(
        wb,
        "Architecture Overview",
        [
            "Component",
            "Role",
            "Technology",
            "Implementation",
            "Notes",
        ],
        [
            [
                "Frontend",
                "SPA for requirements, diagrams, evaluation UX",
                "Vite + React (planned)",
                "PLANNED",
                "CORS expects localhost:5173; static host via S3+CloudFront when deployed",
            ],
            [
                "API",
                "REST API /api/v1 + /health",
                "FastAPI",
                "IMPLEMENTED",
                "backend/app/main.py",
            ],
            [
                "Lambda adapter",
                "Serverless HTTP entry",
                "Mangum",
                "IMPLEMENTED (code)",
                "app.lambda_handler.handler — deploy wiring PLANNED",
            ],
            [
                "API Gateway",
                "Public HTTPS front for Lambda",
                "HTTP API",
                "PLANNED",
                "See docs/DEPLOYMENT.md",
            ],
            [
                "Requirements Interpreter",
                "Normalize user requirements",
                "agent/requirements_interpreter.py",
                "IMPLEMENTED",
                "May call Bedrock or local heuristics",
            ],
            [
                "Architecture Planner",
                "Produce Architecture JSON",
                "agent + Bedrock",
                "IMPLEMENTED",
                "Never emits Terraform/CloudFormation",
            ],
            [
                "Diagram Builder/Parser",
                "Architecture ↔ Draw.io XML",
                "drawio_builder / drawio_parser",
                "IMPLEMENTED",
                "Uses defusedxml for safe parse",
            ],
            [
                "Rules Engine",
                "Deterministic evaluation",
                "rules/*",
                "IMPLEMENTED",
                "security, cost, reliability, scalability, best practices",
            ],
            [
                "Gap Analyst",
                "LLM gap narrative",
                "agent/gap_analyst.py",
                "IMPLEMENTED",
                "Falls back in DEMO_MODE",
            ],
            [
                "Optimization Planner",
                "Prioritized recommendations",
                "agent/optimization_planner.py",
                "IMPLEMENTED",
                "Qualitative cost notes only",
            ],
            [
                "Report/Export",
                "JSON / Draw.io / XLSX",
                "exporter + report_generator",
                "IMPLEMENTED",
                "No live price quotes",
            ],
            [
                "Local demo path",
                "Run without AWS",
                "DEMO_MODE + local data/",
                "IMPLEMENTED",
                "Mock Bedrock when credentials absent",
            ],
        ],
    )

    add_sheet(
        wb,
        "AWS Service Inventory",
        [
            "AWS Service",
            "Purpose in InfraPilot",
            "Implementation",
            "Config / env",
            "Cost posture (qualitative)",
            "Notes",
        ],
        [
            [
                "Amazon Bedrock",
                "LLM for plan/gap/optimize",
                "IMPLEMENTED client",
                "BEDROCK_MODEL_ID, BEDROCK_REGION, AWS_REGION",
                "Primary variable cost driver — use Haiku-class models; set DEMO_MODE for free local",
                "Requires model access in account; no prices invented here",
            ],
            [
                "AWS Lambda",
                "Host FastAPI via Mangum",
                "Code IMPLEMENTED; function PLANNED",
                "Handler app.lambda_handler.handler",
                "Prefer low memory/timeout for demos",
                "Package size: include agent/, rules/, catalog/",
            ],
            [
                "API Gateway (HTTP API)",
                "HTTPS invoke Lambda",
                "PLANNED",
                "CORS origins from CORS_ORIGINS",
                "Low for light traffic",
                "Prefer HTTP API over REST for cost",
            ],
            [
                "Amazon S3 (artifacts)",
                "Uploads / exports",
                "Service IMPLEMENTED; bucket PLANNED",
                "S3_BUCKET",
                "Storage + requests — keep lifecycle short for demos",
                "Separate from frontend bucket recommended",
            ],
            [
                "Amazon S3 (frontend)",
                "Static Vite dist",
                "PLANNED",
                "See DEPLOYMENT.md",
                "Negligible for static assets",
                "Block public access; use OAC with CloudFront",
            ],
            [
                "Amazon CloudFront",
                "CDN for frontend",
                "PLANNED",
                "Distribution + OAC",
                "Low for low traffic; watch invalidations",
                "Optional for pure API demos",
            ],
            [
                "Amazon DynamoDB",
                "Project metadata",
                "Service IMPLEMENTED; table PLANNED",
                "DYNAMODB_TABLE",
                "Prefer on-demand billing for demos",
                "Key: id; payload in data attribute",
            ],
            [
                "Amazon CloudWatch",
                "Lambda logs/metrics",
                "PLANNED ops",
                "Log group auto with Lambda",
                "Set retention (e.g. 7–14 days) to limit cost",
                "No custom dashboards required for MVP",
            ],
            [
                "IAM",
                "Execution + deploy roles",
                "PLANNED",
                "OIDC deploy role placeholders",
                "N/A",
                "Least privilege templates in DEPLOYMENT.md — not production-certified policies",
            ],
            [
                "GitHub Actions OIDC",
                "Keyless deploy",
                "SCAFFOLDED workflow",
                "AWS_ROLE_TO_ASSUME placeholder",
                "N/A",
                "No long-lived access keys in CI",
            ],
        ],
    )

    add_sheet(
        wb,
        "Data Flow",
        [
            "Step",
            "Input",
            "Transform",
            "Output",
            "Storage",
            "Implementation",
        ],
        [
            [
                "1",
                "User requirements / clarifications",
                "Requirements Interpreter",
                "Requirements JSON",
                "Project record",
                "IMPLEMENTED",
            ],
            [
                "2",
                "Requirements JSON",
                "Architecture Planner (+ Bedrock or mock)",
                "Architecture JSON (nodes/edges)",
                "Project + optional S3",
                "IMPLEMENTED",
            ],
            [
                "3",
                "Architecture JSON",
                "Diagram Builder",
                "Draw.io XML",
                "S3 or local files",
                "IMPLEMENTED",
            ],
            [
                "4",
                "Uploaded Draw.io / JSON",
                "Parser → Architecture",
                "Architecture JSON",
                "Project",
                "IMPLEMENTED",
            ],
            [
                "5",
                "Architecture (+ Requirements)",
                "Rules Engine",
                "Evaluation findings + ScoreReport",
                "Project / export",
                "IMPLEMENTED",
            ],
            [
                "6",
                "Findings + architecture",
                "Gap Analyst",
                "Narrative gap items",
                "Export",
                "IMPLEMENTED",
            ],
            [
                "7",
                "Findings + scores",
                "Optimization Planner",
                "Prioritized recommendations",
                "Export",
                "IMPLEMENTED",
            ],
            [
                "8",
                "Combined artifacts",
                "Report / Exporter",
                "JSON / Draw.io / XLSX downloads",
                "Client download; optional S3",
                "IMPLEMENTED",
            ],
        ],
    )

    add_sheet(
        wb,
        "Validation Rules",
        [
            "Category",
            "Module",
            "What it checks (summary)",
            "Statuses",
            "Scoring weight (default)",
            "Implementation",
        ],
        [
            [
                "Security",
                "rules/security_rules.py",
                "Exposure, encryption cues, auth patterns at design level",
                "Pass/Fail/Warning/Unknown/N/A",
                "0.30",
                "IMPLEMENTED",
            ],
            [
                "Cost",
                "rules/cost_rules.py",
                "Qualitative cost smells (always-on, oversize) — no AWS price quotes",
                "Pass/Fail/Warning/Unknown/N/A",
                "0.15",
                "IMPLEMENTED",
            ],
            [
                "Reliability",
                "rules/reliability_rules.py",
                "Multi-AZ cues, backups, single points of failure",
                "Pass/Fail/Warning/Unknown/N/A",
                "0.25",
                "IMPLEMENTED",
            ],
            [
                "Scalability",
                "rules/scalability_rules.py",
                "Autoscaling / serverless fit vs traffic assumption",
                "Pass/Fail/Warning/Unknown/N/A",
                "0.15",
                "IMPLEMENTED",
            ],
            [
                "Best practices",
                "rules/best_practice_rules.py",
                "Catalog alignment, naming, Well-Architected-ish heuristics",
                "Pass/Fail/Warning/Unknown/N/A",
                "0.15",
                "IMPLEMENTED",
            ],
            [
                "Engine",
                "rules/engine.py",
                "Aggregates findings into category + overall 0–100 scores",
                "ScoreReport",
                "N/A",
                "IMPLEMENTED",
            ],
        ],
    )

    add_sheet(
        wb,
        "Security and IAM",
        [
            "Control area",
            "Guidance",
            "Implementation",
            "Placeholder / action",
        ],
        [
            [
                "Secrets in repo",
                "Never commit .env, keys, or tokens; use .env.example placeholders only",
                "DOCUMENTED + .gitignore",
                "Rotate if anything leaked",
            ],
            [
                "Lambda execution role",
                "Least privilege: bedrock:InvokeModel (scoped), s3:Get/Put on artifacts prefix, dynamodb CRUD on one table, logs",
                "PLANNED",
                "REPLACE_ACCOUNT_ID / resource ARNs — do not paste broad AdministratorAccess",
            ],
            [
                "Deploy role (OIDC)",
                "s3:PutObject on frontend+artifacts buckets, lambda:UpdateFunctionCode, cloudfront:CreateInvalidation; deny * otherwise",
                "SCAFFOLDED in deploy.yml",
                "Set AWS_ROLE_TO_ASSUME GitHub variable",
            ],
            [
                "Bedrock",
                "Prefer small models; maxTokens capped in client; DEMO_MODE for offline",
                "IMPLEMENTED safeguards (partial)",
                "Enable only needed models in Bedrock console",
            ],
            [
                "CORS",
                "Restrict CORS_ORIGINS to known frontends",
                "IMPLEMENTED setting",
                "Update for CloudFront domain when live",
            ],
            [
                "XML uploads",
                "Parse with defusedxml; size limits recommended at API Gateway/Lambda",
                "IMPLEMENTED parser choice",
                "Configure payload limits in gateway (PLANNED)",
            ],
            [
                "Public access",
                "Frontend via CloudFront OAC; block S3 public ACLs",
                "PLANNED",
                "Follow DEPLOYMENT.md checklist",
            ],
            [
                "Production IAM policies",
                "This workbook does NOT invent production-ready IAM JSON",
                "OUT OF SCOPE here",
                "Review with security owner before prod",
            ],
        ],
    )

    add_sheet(
        wb,
        "Risks and Recommendations",
        [
            "Risk",
            "Likelihood",
            "Impact",
            "Mitigation",
            "Status",
        ],
        [
            [
                "Bedrock cost overrun from chatty UX",
                "Medium",
                "Medium",
                "DEMO_MODE default locally; Haiku; rate limits; cache sample flows",
                "OPEN",
            ],
            [
                "Frontend not ready while API is",
                "High (current)",
                "Low",
                "Use /docs OpenAPI + demo endpoints; ship SPA incrementally",
                "ACCEPTED",
            ],
            [
                "Users treat scores as compliance certification",
                "Medium",
                "High",
                "Document limitations; label SAMPLE data; no deploy/IaC generation",
                "DOCUMENTED",
            ],
            [
                "Overly broad IAM during setup",
                "Medium",
                "High",
                "Use least-privilege templates; OIDC; no long-lived CI keys",
                "OPEN",
            ],
            [
                "Lambda package missing agent/rules/catalog",
                "Medium",
                "High",
                "Document packaging paths in DEPLOYMENT.md; CI build step",
                "OPEN",
            ],
            [
                "Draw.io malicious XML",
                "Low",
                "Medium",
                "defusedxml + size limits + content-type checks",
                "PARTIAL",
            ],
        ],
    )

    add_sheet(
        wb,
        "CI-CD and Operations",
        [
            "Area",
            "Tooling",
            "Trigger",
            "Implementation",
            "Notes",
        ],
        [
            [
                "Lint",
                "ruff",
                "PR / push (ci.yml)",
                "SCAFFOLDED",
                "backend + agent + rules",
            ],
            [
                "Unit/integration tests",
                "pytest",
                "PR / push",
                "IMPLEMENTED tests; CI SCAFFOLDED",
                "DEMO_MODE friendly",
            ],
            [
                "Frontend build",
                "npm ci && npm run build",
                "CI when frontend/ exists",
                "SCAFFOLDED (conditional)",
                "Vite dist artifact",
            ],
            [
                "Backend package",
                "zip Lambda bundle",
                "CI/CD",
                "SCAFFOLDED",
                "Include repo-root packages",
            ],
            [
                "Deploy",
                "GitHub Actions OIDC → AWS",
                "workflow_dispatch",
                "SCAFFOLDED placeholders",
                "Manual approve deploy",
            ],
            [
                "Health check",
                "GET /health",
                "Post-deploy smoke",
                "IMPLEMENTED endpoint",
                "Returns demo_mode flag",
            ],
            [
                "Observability",
                "CloudWatch Logs",
                "Runtime",
                "PLANNED",
                "Set log retention",
            ],
            [
                "Teardown",
                "AWS CLI delete resources",
                "Manual",
                "DOCUMENTED",
                "docs/DEPLOYMENT.md",
            ],
        ],
    )

    add_sheet(
        wb,
        "Cost Assumptions",
        [
            "Item",
            "Assumption",
            "Guidance",
            "Numeric price",
            "Implementation note",
        ],
        [
            [
                "Bedrock invocations",
                "Dominant cost for non-demo usage",
                "Use Claude Haiku-class; disable when DEMO_MODE=true",
                "PLACEHOLDER — look up current Bedrock pricing in AWS console/docs",
                "Client caps maxTokens; no app-side price estimates",
            ],
            [
                "Lambda",
                "Low request volume for demos",
                "128–512 MB; short timeout; avoid provisioned concurrency",
                "PLACEHOLDER — use AWS Pricing Calculator",
                "Prefer free-tier eligible configs where available",
            ],
            [
                "API Gateway",
                "HTTP API, light traffic",
                "Prefer HTTP API over REST API",
                "PLACEHOLDER",
                "PLANNED resource",
            ],
            [
                "DynamoDB",
                "On-demand, few projects",
                "On-demand billing for demos",
                "PLACEHOLDER",
                "Avoid unused provisioned capacity",
            ],
            [
                "S3 + CloudFront",
                "Static SPA + small artifacts",
                "Lifecycle expire uploads; minimize invalidations",
                "PLACEHOLDER",
                "Separate buckets recommended",
            ],
            [
                "CloudWatch Logs",
                "Retention drives cost",
                "7–14 day retention for non-prod",
                "PLACEHOLDER",
                "Set retention on first deploy",
            ],
            [
                "Local development",
                "Zero AWS spend",
                "DEMO_MODE=true without credentials",
                "N/A ($0 AWS)",
                "IMPLEMENTED path",
            ],
            [
                "Disclaimer",
                "This sheet never invents production prices",
                "Replace PLACEHOLDER with values from AWS Pricing at time of estimate",
                "NOT PROVIDED",
                "InfraPilot cost rules are qualitative only",
            ],
        ],
    )

    # Cover note sheet first
    cover = wb.create_sheet("README", 0)
    cover["A1"] = "InfraPilot AI — Architecture Workbook"
    cover["A1"].font = Font(bold=True, size=14)
    cover["A3"] = (
        "Status labels: IMPLEMENTED = present in repository code; "
        "PLANNED / SCAFFOLDED = documented or workflow-only; "
        "PLACEHOLDER = fill with your account values. "
        "This workbook does not contain production IAM policy JSON or AWS list prices."
    )
    cover["A3"].alignment = Alignment(wrap_text=True)
    cover.merge_cells("A3:D6")
    cover["A8"] = "Related diagrams: docs/architecture/drawio/*.drawio"
    cover["A9"] = "Related docs: docs/DEPLOYMENT.md, LOCAL_DEVELOPMENT.md, SECURITY.md, LIMITATIONS.md"
    cover.column_dimensions["A"].width = 100

    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
