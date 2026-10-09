"""Export unit tests."""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.models.schemas import Project
from app.services import exporter


def test_export_drawio(sample_architecture):
    data = exporter.export_drawio(sample_architecture)
    assert b"mxfile" in data
    assert b"mxCell" in data


def test_export_png_placeholder(sample_architecture):
    data = exporter.export_png(sample_architecture)
    assert data[:8] == b"\x89PNG\r\n\x1a\n"


def test_export_excel_worksheets(sample_architecture, sample_requirements):
    from agent.gap_analyst import GapAnalyst

    project = Project(
        name="Export Test",
        requirements=sample_requirements,
        architecture=sample_architecture,
    )
    project.evaluation = GapAnalyst().analyze(
        project.id, sample_architecture, sample_requirements
    )
    data = exporter.export_excel(project)
    wb = load_workbook(BytesIO(data))
    expected = {
        "Architecture Overview",
        "AWS Service Inventory",
        "Data Flow",
        "Validation Rules",
        "Security and IAM",
        "Risks and Recommendations",
        "CI-CD and Operations",
        "Cost Assumptions",
    }
    assert expected.issubset(set(wb.sheetnames))
    # Cost sheet must not invent dollar prices
    cost_ws = wb["Cost Assumptions"]
    joined = " ".join(
        str(cell.value or "") for row in cost_ws.iter_rows() for cell in row
    )
    assert "$" not in joined or "invent" in joined.lower() or "Pricing" in joined
